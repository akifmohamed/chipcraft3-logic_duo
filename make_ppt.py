#!/usr/bin/env python3
# Round-0 deck generator for Team LOGIC DUO — ChipCraft 3.0
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY   = RGBColor(0x0F, 0x20, 0x41)
TEAL   = RGBColor(0x00, 0x96, 0x88)
AMBER  = RGBColor(0xFF, 0xB3, 0x00)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
DARK   = RGBColor(0x21, 0x21, 0x21)
GRAY   = RGBColor(0x55, 0x55, 0x55)
LIGHT  = RGBColor(0xF2, 0xF6, 0xF8)
ROW_A  = RGBColor(0xE8, 0xF2, 0xF1)
ROW_B  = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height

def rect(slide, x, y, w, h, color):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh

def text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0):
    """runs = list of paragraphs; each = (text, size, bold, color, [space_after])"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, (t, size, bold, color, *rest) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        if rest:
            p.space_after = Pt(rest[0])
        r = p.add_run(); r.text = t
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
        r.font.name = "Calibri"
    return tb

def header(slide, num, title):
    rect(slide, 0, 0, SW, Inches(1.05), NAVY)
    rect(slide, 0, Inches(1.05), SW, Inches(0.06), TEAL)
    text(slide, Inches(0.5), Inches(0.12), Inches(1.5), Inches(0.8),
         [(f"SLIDE {num}", 12, True, AMBER, 0),
          ("ROUND 0", 10, False, WHITE, 0)], anchor=MSO_ANCHOR.MIDDLE)
    text(slide, Inches(2.2), Inches(0.1), Inches(10.6), Inches(0.85),
         [(title, 30, True, WHITE, 0)], anchor=MSO_ANCHOR.MIDDLE)

def notes(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt

# ================= SLIDE 1 : TITLE =================
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, SW, SH, NAVY)
rect(s, 0, Inches(4.35), SW, Inches(0.07), TEAL)
text(s, Inches(0.8), Inches(0.55), Inches(11.7), Inches(0.5),
     [("CHIPCRAFT 3.0  —  RTL TO PHYSICAL DESIGN HACKATHON", 16, True, AMBER, 0)],
     align=PP_ALIGN.CENTER)
text(s, Inches(0.8), Inches(1.45), Inches(11.7), Inches(1.5),
     [("SMART WATER", 54, True, WHITE, 4),
      ("DISTRIBUTION SCHEDULER", 54, True, WHITE, 0)],
     align=PP_ALIGN.CENTER)
text(s, Inches(0.8), Inches(3.5), Inches(11.7), Inches(0.6),
     [("Priority-Aware · Emergency-Ready · Conflict-Free Valve Control", 20, False, TEAL, 0)],
     align=PP_ALIGN.CENTER)
text(s, Inches(0.8), Inches(4.75), Inches(11.7), Inches(2.3),
     [("TEAM LOGIC DUO   |   Team #32", 22, True, WHITE, 10),
      ("Akif Mohamed  (Team Lead)   ·   [ add teammate name ]", 16, False, WHITE, 4),
      ("Government College of Engineering, Srirangam", 16, False, WHITE, 14),
      ("Flow:  Verilog RTL  →  VCS Simulation  →  DC Synthesis  →  ICC2 Layout   (SAED 32nm)",
       13, False, AMBER, 0)],
     align=PP_ALIGN.CENTER)
notes(s, "Good morning! We are TEAM LOGIC DUO from Government College of Engineering, "
         "Srirangam. Our problem is the Smart Water Distribution Scheduler. "
         "[Say your teammate's name too — and remember to type it on this slide first!]")

# ================= SLIDE 2 : PS IN 2 LINES =================
s = prs.slides.add_slide(BLANK)
header(s, 2, "PROBLEM STATEMENT")
rect(s, Inches(0.9), Inches(1.7), Inches(11.5), Inches(3.2), LIGHT)
rect(s, Inches(0.9), Inches(1.7), Inches(0.12), Inches(3.2), TEAL)
text(s, Inches(1.35), Inches(2.0), Inches(10.8), Inches(2.7),
     [("Design a controller that schedules water to multiple zones", 26, True, NAVY, 10),
      ("based on DEMAND, PRIORITY and AVAILABLE SUPPLY.", 26, True, NAVY, 22),
      ("It serves the most urgent zone first, opens only ONE valve at a time,", 20, False, DARK, 8),
      ("and never wastes water when supply is short.", 20, False, DARK, 0)],
     spacing=1.15)
text(s, Inches(0.9), Inches(5.35), Inches(11.5), Inches(1.6),
     [("Why it matters:", 16, True, TEAL, 6),
      ("In real cities, low-pressure zones wait while high-pressure zones overflow.", 15, False, GRAY, 4),
      ("Our chip makes distribution FAIR, SAFE and ACCOUNTABLE.", 15, False, GRAY, 0)])
notes(s, "The PS in two lines: Design a controller that schedules water to multiple zones "
         "based on demand, priority and available supply — serving the most urgent zone first, "
         "opening only one valve at a time, and never wasting water when supply is short. "
         "That's it. Simple, but every objective in the PS has to be covered.")

# ================= SLIDE 3 : INNOVATION / SOLUTION =================
s = prs.slides.add_slide(BLANK)
header(s, 3, "OUR INNOVATIVE SOLUTION")

def bullet_card(slide, x, y, w, h, ic, title_, body_):
    rect(slide, x, y, w, h, LIGHT)
    rect(slide, x, y, Inches(0.1), h, TEAL)
    text(slide, x + Inches(0.28), y + Inches(0.12), w - Inches(0.45), h - Inches(0.2),
         [(ic + "  " + title_, 16, True, NAVY, 5),
          (body_, 13.5, False, DARK, 0)], spacing=1.05)

bullet_card(s, Inches(0.55), Inches(1.5),  Inches(6.1), Inches(1.55), "🏆",
            "Priority + Emergency Arbiter",
            "4 zones, priority 0-3. Priority 3 = EMERGENCY jumps the queue.\nTie-break: lowest zone index — fair and deterministic.")
bullet_card(s, Inches(6.85), Inches(1.5),  Inches(6.0), Inches(1.55), "🔒",
            "Conflict-Free One-Hot Valves",
            "Exactly one valve can open at a time (one-hot by design).\nA running flush is never aborted — valve safety first.")
bullet_card(s, Inches(0.55), Inches(3.25), Inches(6.1), Inches(1.55), "🧠",
            "SMART Skip-to-Fit Scheduling  (our innovation)",
            "If the top zone's demand > available water, we serve a smaller\nzone that FITS instead of stalling the whole city. No starvation!")
bullet_card(s, Inches(6.85), Inches(3.25), Inches(6.0), Inches(1.55), "📊",
            "Accountability Built-In",
            "Total + per-zone cycle counters and a LIVE insufficient-supply flag.\nEvery litre served is tracked — audit-ready.")
rect(s, Inches(0.55), Inches(5.1), Inches(12.3), Inches(1.7), NAVY)
text(s, Inches(0.85), Inches(5.25), Inches(11.7), Inches(1.4),
     [("Design: FSM (IDLE → SERVE → WAIT/DONE)  +  demand tracker  +  priority arbiter  +  one-hot valve driver  +  counters",
       15, True, WHITE, 8),
      ("Plan: clean synthesizable Verilog for 32nm, covered by a self-checking testbench (7 test groups) across the full RTL→GDS flow",
       14, False, AMBER, 0)], anchor=MSO_ANCHOR.MIDDLE, spacing=1.1)
notes(s, "Our solution has four building blocks. Two are required by the PS — the priority "
         "arbiter with emergency support, and the one-hot valve safety. Two are OUR innovation: "
         "first, skip-to-fit scheduling — if the biggest emergency zone needs 200 units but only "
         "100 are available, we don't stall the city; we serve a smaller zone that fits and flag "
         "the shortage. Second, accountability — total and per-zone counters plus a live "
         "insufficient-supply flag. It's a simple FSM — tiny and fast for 32nm, and the whole "
         "plan maps cleanly onto the RTL-to-GDS flow.")

# ================= SLIDE 4 : I/O + PARAMETERS =================
s = prs.slides.add_slide(BLANK)
header(s, 4, "INTERFACE — INPUTS, OUTPUTS & PARAMETERS")

def mini_table(slide, x, y, w, rows, colw, title_, fs=11.5):
    n = len(rows)
    tb = slide.shapes.add_table(n + 1, 2, x, y, w, Inches(0.32 * (n + 1))).table
    tb.columns[0].width = colw
    tb.columns[1].width = w - colw
    hdr = [("SIGNAL", "WHAT IT IS")][0]
    for j, htxt in enumerate(hdr):
        c = tb.cell(0, j); c.text = htxt
        c.fill.solid(); c.fill.fore_color.rgb = NAVY
        p = c.text_frame.paragraphs[0]; p.font.size = Pt(11); p.font.bold = True
        p.font.color.rgb = WHITE; p.font.name = "Calibri"
    for i, (a, b) in enumerate(rows):
        for j, v in enumerate((a, b)):
            c = tb.cell(i + 1, j); c.text = v
            c.fill.solid(); c.fill.fore_color.rgb = ROW_A if i % 2 == 0 else ROW_B
            p = c.text_frame.paragraphs[0]
            p.font.size = Pt(fs); p.font.name = "Calibri"
            p.font.color.rgb = DARK
            p.font.bold = (j == 0)
    text(slide, x, y - Inches(0.38), w, Inches(0.35),
         [(title_, 14, True, TEAL, 0)])
    return tb

IN_ROWS = [
    ("clk, rst_n", "100 MHz clock, async reset"),
    ("zone_req[3:0]", "zone demands water (one bit/zone)"),
    ("zone_prio[7:0]", "2-bit priority / zone (3 = EMERGENCY)"),
    ("zone_demand[31:0]", "8-bit water units wanted per zone"),
    ("supply_avail[7:0]", "water units available now"),
    ("supply_valid", "pump ON"),
]
OUT_ROWS = [
    ("valve_open[3:0]", "one-hot valve commands (safe!)"),
    ("grant_ack[3:0]", "pulse: that zone is served"),
    ("cycles_done[31:0]", "total completed cycles"),
    ("zone_cycles[63:0]", "16-bit cycle count per zone"),
    ("insufficient", "1 = demand but not enough water"),
    ("busy, status[2:0]", "serve in progress; state code"),
]
PAR_ROWS = [
    ("NZONES = 4", "number of zones"),
    ("DATA_W = 8", "water unit width (0–255)"),
    ("SERVE_CYCLES = 16", "clocks per distribution"),
    ("Clock = 100 MHz", "10 ns, SAED 32nm"),
]
mini_table(s, Inches(0.55), Inches(1.85), Inches(6.1),  IN_ROWS,  Inches(2.3), "INPUTS")
mini_table(s, Inches(6.85), Inches(1.85), Inches(6.0),  OUT_ROWS, Inches(2.3), "OUTPUTS")
mini_table(s, Inches(0.55), Inches(5.05), Inches(6.1),  PAR_ROWS, Inches(2.5), "KEY PARAMETERS & VALUES")
text(s, Inches(6.85), Inches(4.7), Inches(6.0), Inches(2.4),
     [("STATUS CODES", 14, True, TEAL, 6),
      ("0 = IDLE (nothing to do)", 13, False, DARK, 3),
      ("1 = SERVE (valve open, water flowing)", 13, False, DARK, 3),
      ("2 = WAIT (insufficient supply)", 13, False, DARK, 3),
      ("3 = CYCLE DONE (bookkeeping pulse)", 13, False, DARK, 8),
      ("PRIORITY VALUES", 14, True, TEAL, 6),
      ("0 = Low · 1 = Medium · 2 = High · 3 = Emergency", 13, False, DARK, 0)])
notes(s, "Our interface. Inputs: each zone sends a request, a 2-bit priority and its demand in "
         "water units. The plant sends available supply and pump status. Outputs: one-hot valve "
         "commands — the safety core of our design — an ack pulse when a zone is served, total "
         "and per-zone cycle counters, and a live insufficient-supply flag. Everything is "
         "parameterized: 4 zones, 8-bit demands, 16-clock serve time, 100 MHz target.")

# ================= SLIDE 5 : EXPECTED OUTCOMES / VALUES =================
s = prs.slides.add_slide(BLANK)
header(s, 5, "EXPECTED OUTCOMES & KEY VALUES")

def check_row(slide, x, y, w, label, value):
    rect(slide, x, y, w, Inches(0.52), LIGHT)
    text(slide, x + Inches(0.15), y + Inches(0.03), w * 0.62, Inches(0.45),
         [("→  " + label, 13.5, True, DARK, 0)], anchor=MSO_ANCHOR.MIDDLE)
    text(slide, x + w * 0.6, y + Inches(0.03), w * 0.38, Inches(0.45),
         [(value, 13, False, TEAL, 0)], anchor=MSO_ANCHOR.MIDDLE)

L = [("PS objectives to be covered", "target 6 / 6"),
     ("Self-checking TB (7 test groups)", "expected ALL PASS"),
     ("Valve conflicts (one-hot)", "expected 0 conflicts"),
     ("Priority + emergency behaviour", "served first, by design"),
     ("Insufficient-supply handling", "flag + WAIT + skip-to-fit"),
     ("Cycle tracking accuracy", "total = Σ per-zone")]
R = [("Clock target (SAED 32nm)", "100 MHz / 10 ns"),
     ("Logic style", "Synthesizable Verilog, no latches"),
     ("Synthesis (DC)", "expect: timing met + clean QoR"),
     ("Physical design (ICC2)", "expect: P&R clean + GDS"),
     ("Final deliverables", "RTL+TB+SDC+Netlist+Rpts+GDS"),
     ("Design area class", "4-zone controller, 32-bit counters")]
y0 = Inches(1.55)
for i, (a, b) in enumerate(L):
    check_row(s, Inches(0.55), y0 + Inches(0.6) * i, Inches(6.1), a, b)
for i, (a, b) in enumerate(R):
    check_row(s, Inches(6.85), y0 + Inches(0.6) * i, Inches(6.0), a, b)

rect(s, Inches(0.55), Inches(5.5), Inches(12.3), Inches(1.35), NAVY)
text(s, Inches(0.85), Inches(5.6), Inches(11.7), Inches(1.15),
     [("EXPECTED BOTTOM LINE:", 15, True, AMBER, 6),
      ("A fair, safe and auditable water scheduler — all 6 objectives met, RTL verified, "
       "timing-closed 32nm layout at 100 MHz.", 15, False, WHITE, 0)],
     anchor=MSO_ANCHOR.MIDDLE, spacing=1.1)
notes(s, "What we EXPECT to deliver. All six PS objectives covered. A self-checking "
         "testbench with 7 test groups, targeting all pass. By design the valves are one-hot, "
         "so we expect zero conflicts. The counters will be exact: total equals the sum of "
         "per-zone counts. Technical targets: 100 MHz in SAED 32nm, latch-free synthesizable "
         "Verilog, timing-closed synthesis in DC and clean place-and-route with GDS in ICC2. "
         "Expected bottom line: fair, safe, auditable.")

prs.save("/home/user/chipcraft/docs/Round0_LOGIC_DUO.pptx")
print("SAVED: docs/Round0_LOGIC_DUO.pptx  (5 slides)")
